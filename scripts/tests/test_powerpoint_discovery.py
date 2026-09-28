"""Opt-in, no-model PowerPoint discovery and saved-output regressions."""

import asyncio
import importlib.util
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
import uuid
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
spec = importlib.util.spec_from_file_location(
    "powerpoint_benchmark", Path(__file__).resolve().parents[1] / "benchmark-real-apps.py")
benchmark = importlib.util.module_from_spec(spec)
spec.loader.exec_module(benchmark)


@unittest.skipUnless(
    os.environ.get("MCP_TEST_DESKTOP_INPUT") == "1" and os.environ.get("MCP_TEST_POWERPOINT_SERVER"),
    "Requires an available desktop, no existing PowerPoint process, and MCP_TEST_POWERPOINT_SERVER",
)
class PowerPointDiscoveryTests(unittest.IsolatedAsyncioTestCase):
    async def cli(self, *arguments):
        process = await asyncio.create_subprocess_exec(
            os.environ["MCP_TEST_POWERPOINT_CLI"], *arguments,
            stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE)
        try:
            stdout, stderr = await asyncio.wait_for(process.communicate(), 45)
            self.assertEqual(0, process.returncode, stderr.decode())
            return json.loads(stdout)
        finally:
            if process.returncode is None:
                process.kill()
                await process.wait()

    async def exercise(self, initial_read, edit=False):
        from mcp import ClientSession
        from mcp.client.stdio import StdioServerParameters, stdio_client
        from pptx import Presentation

        with tempfile.TemporaryDirectory(prefix="powerpoint-discovery-") as temporary:
            directory = Path(temporary)
            source, _ = benchmark.create_fixture("powerpoint", directory)
            source = source.rename(source.with_stem(f"input-{uuid.uuid4().hex}"))
            original = source.read_bytes()
            async with stdio_client(StdioServerParameters(
                command=os.environ["MCP_TEST_POWERPOINT_SERVER"], env=dict(os.environ),
            )) as streams:
                async with ClientSession(*streams, read_timeout_seconds=120) as client:
                    await client.initialize()
                    owned = benchmark.OwnedApp("powerpoint", source, directory)
                    try:
                        async def call(tool_name, **arguments):
                            result = await client.call_tool(tool_name, arguments)
                            self.assertFalse(result.is_error, str(result.content))
                            data = json.loads(result.content[0].text)
                            self.assertTrue(data.get("success"), data)
                            return data

                        await call("window_management", action="activate", handle=owned.root)
                        if initial_read == "cli":
                            title = await self.cli(
                                "ui", "find", "--window", owned.root, "--name", "Title TextBox",
                                "--control-type", "Image", "--require-unique")
                            self.assertTrue(title["success"], title)
                            self.assertEqual(1, len(title["items"]))
                            text = await self.cli(
                                "ui", "read", "--window", owned.root, "--element-id", title["items"][0]["id"])
                            self.assertEqual("Cedar kickoff", text["text"].strip())
                            self.assertEqual(original, source.read_bytes())
                            return
                        if initial_read == "snapshot":
                            snapshot = await call("ui_snapshot", windowHandle=owned.root, maxDepth=20)
                            slides = [node for node in benchmark.flatten_tree(snapshot["tree"])
                                      if node.get("type") == "ListItem" and node.get("name") == "Slide Cedar kickoff"]
                            self.assertEqual(1, len(slides), "The first snapshot must expose the slide, without prior input.")
                        elif initial_read == "read":
                            text = await call("ui_read", windowHandle=owned.root, includeChildren=True)
                            self.assertIn("Cedar kickoff", text["text"])
                        title = await call("ui_find", windowHandle=owned.root, name="Title TextBox",
                                           controlType="Image", requireUnique=True, timeoutMs=1000)
                        element = title["items"][0]["id"]
                        text = await call("ui_read", windowHandle=owned.root, elementId=element)
                        self.assertEqual("Cedar kickoff", text["text"].strip())
                        if edit:
                            await call("ui_click", windowHandle=owned.root, elementId=element, doubleClick=True)
                            await call("keyboard_control", windowHandle=owned.root, action="press", key="a", modifiers="ctrl")
                            await call("keyboard_control", windowHandle=owned.root, action="type", text="Cedar discovery verified")
                            destination = directory / "completed.pptx"
                            await call("file_save", windowHandle=owned.root, filePath=str(destination), triggerMode="save_as")
                            for _ in range(100):
                                if destination.exists() and zipfile.is_zipfile(destination):
                                    break
                                await asyncio.sleep(0.1)
                            self.assertTrue(destination.exists() and zipfile.is_zipfile(destination))
                            before = Presentation(source)
                            after = Presentation(destination)
                            expected = [[shape.text for shape in slide.shapes if shape.has_text_frame]
                                        for slide in before.slides]
                            expected[0][0] = "Cedar discovery verified"
                            actual = [[shape.text for shape in slide.shapes if shape.has_text_frame]
                                      for slide in after.slides]
                            self.assertEqual(expected, actual)
                        self.assertEqual(original, source.read_bytes())
                    finally:
                        try:
                            if initial_read == "cli":
                                await self.cli("service", "stop")
                        finally:
                            owned.close()

    async def test_first_snapshot_exposes_slide_and_readable_title(self):
        await self.exercise("snapshot")

    async def test_first_find_exposes_readable_title_without_snapshot(self):
        await self.exercise("find")

    async def test_first_window_read_includes_slide_content(self):
        await self.exercise("read")

    @unittest.skipUnless(os.environ.get("MCP_TEST_POWERPOINT_CLI"), "Requires MCP_TEST_POWERPOINT_CLI")
    async def test_cli_first_find_and_separate_read_expose_title(self):
        await self.exercise("cli")

    async def test_discovered_title_edit_is_verified_in_saved_presentation(self):
        await self.exercise("snapshot", edit=True)


if __name__ == "__main__":
    unittest.main()
