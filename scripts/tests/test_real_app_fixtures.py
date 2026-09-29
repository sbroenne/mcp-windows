import json
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import real_app_fixtures as fixtures


class RealAppFixtureTests(unittest.TestCase):
    def test_notepad_requires_all_edits_and_preserved_content(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = fixtures.create_fixture("notepad", Path(directory))
            self.assertFalse(fixtures.verify_output("notepad", output)["success"])
            output.write_text(source.read_text(encoding="utf-8"), encoding="utf-8")
            self.assertFalse(fixtures.verify_output("notepad", output)["success"])
            output.write_text(fixtures.NOTEPAD_EXPECTED, encoding="utf-8")
            self.assertTrue(fixtures.verify_output("notepad", output)["success"])
            output.write_text(fixtures.NOTEPAD_EXPECTED + "Extra text", encoding="utf-8")
            self.assertFalse(fixtures.verify_output("notepad", output)["success"])

    def test_word_requires_text_heading_style_and_unchanged_paragraphs(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            source, output = fixtures.create_fixture("word", Path(directory))
            document = Document(source)
            document.paragraphs[1].text = fixtures.WORD_EXPECTED[1]
            document.save(output)
            self.assertFalse(fixtures.verify_output("word", output)["success"])
            document.paragraphs[0].style = "Heading 1"
            document.save(output)
            self.assertTrue(fixtures.verify_output("word", output)["success"])
            document.paragraphs[-1].text = "Lost the original paragraph"
            document.save(output)
            self.assertFalse(fixtures.verify_output("word", output)["success"])

    def test_notepad_preserves_trailing_newlines_exactly(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "completed.txt"
            for text, success in (
                (fixtures.NOTEPAD_EXPECTED, True),
                (fixtures.NOTEPAD_EXPECTED.replace("\n", "\r\n"), True),
                (fixtures.NOTEPAD_EXPECTED.rstrip("\n"), False),
                (fixtures.NOTEPAD_EXPECTED + "\n", False),
                (fixtures.NOTEPAD_EXPECTED + "\n\n", False),
            ):
                for encoding in ("utf-8", "utf-8-sig", "utf-16"):
                    with self.subTest(text=repr(text), encoding=encoding):
                        output.write_bytes(text.encode(encoding))
                        self.assertEqual(success, fixtures.verify_output("notepad", output)["success"])

    def test_powerpoint_requires_edit_and_slide_order(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            source, output = fixtures.create_fixture("powerpoint", Path(directory))
            deck = Presentation(source)
            deck.slides[0].shapes.title.text = "Cedar launch briefing"
            deck.slides[1].placeholders[1].text_frame.paragraphs[1].text = "Launch: December"
            deck.save(output)
            self.assertFalse(fixtures.verify_output("powerpoint", output)["success"])
            slides = deck.slides._sldIdLst
            slides.insert(1, slides[2])
            deck.save(output)
            self.assertTrue(fixtures.verify_output("powerpoint", output)["success"])

    def test_chrome_requires_submitted_values_not_confirmation_text(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "submission.json"
            output.write_text(json.dumps({"success": True}), encoding="utf-8")
            self.assertFalse(fixtures.verify_output("chrome", output)["success"])
            output.write_text(json.dumps(fixtures.CHROME_EXPECTED), encoding="utf-8")
            self.assertTrue(fixtures.verify_output("chrome", output)["success"])
            wrong = {**fixtures.CHROME_EXPECTED, "attendees": "11"}
            output.write_text(json.dumps(wrong), encoding="utf-8")
            self.assertFalse(fixtures.verify_output("chrome", output)["success"])

    def test_task_prompts_do_not_name_tools_or_supply_input_document_contents(self):
        for app in fixtures.APPS:
            prompt = fixtures.task_prompt(app, Path(r"C:\benchmark\completed"))
            for name in ("ui_", "file_save", "keyboard_control", "screenshot_control", "mouse_control"):
                self.assertNotIn(name, prompt)
        self.assertNotIn("Owner: Maya Chen", fixtures.task_prompt("notepad", Path("completed.txt")))


if __name__ == "__main__":
    unittest.main()
