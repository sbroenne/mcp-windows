"""Task inputs and independent saved-output checks for the real-app benchmark."""

import json
from pathlib import Path
import zipfile


APPS = ("notepad", "word", "powerpoint", "chrome")
NOTEPAD_SOURCE = (
    "Project Cedar launch checklist\n"
    "Owner: Maya Chen\nReview date: 14 October 2026\nStatus: Draft\n\n"
    "Actions\n- Confirm venue\n- Send invitations\n"
)
NOTEPAD_EXPECTED = NOTEPAD_SOURCE.replace("14 October", "21 October").replace(
    "Status: Draft", "Status: Ready for review"
) + "- Book the projector\n"
WORD_SOURCE = [
    "Cedar project update",
    "The launch is planned for 14 October 2026. Maya Chen owns the rollout.",
    "Review meeting",
    "The team will review the budget and confirm the venue.",
]
WORD_EXPECTED = [
    WORD_SOURCE[0], WORD_SOURCE[1].replace("14 October", "21 October"), *WORD_SOURCE[2:]
]
POWERPOINT_SOURCE = [
    ["Cedar kickoff", "Project update"],
    ["Schedule", "Discovery: October", "Launch: November"],
    ["Next steps", "Assign owners", "Confirm budget"],
]
POWERPOINT_EXPECTED = [
    ["Cedar launch briefing", "Project update"],
    POWERPOINT_SOURCE[2],
    ["Schedule", "Discovery: October", "Launch: December"],
]
CHROME_EXPECTED = {
    "name": "Maya Chen",
    "email": "maya.chen@example.test",
    "department": "Operations",
    "attendees": "12",
    "date": "2026-10-21",
    "projector": "yes",
    "notes": "Bring printed agendas.",
}


def create_fixture(app, directory):
    directory.mkdir(parents=True, exist_ok=True)
    extensions = {"notepad": "txt", "word": "docx", "powerpoint": "pptx", "chrome": "json"}
    extension = extensions[app]
    source, output = directory / f"source.{extension}", directory / f"completed.{extension}"
    if app == "notepad":
        source.write_text(NOTEPAD_SOURCE, encoding="utf-8")
    elif app == "word":
        from docx import Document
        document = Document()
        for index, text in enumerate(WORD_SOURCE):
            document.add_paragraph(text, style="Heading 2" if index == 2 else "Normal")
        document.save(source)
    elif app == "powerpoint":
        from pptx import Presentation
        deck = Presentation()
        for texts in POWERPOINT_SOURCE:
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = texts[0]
            body = slide.placeholders[1].text_frame
            body.text = texts[1]
            for text in texts[2:]:
                body.add_paragraph().text = text
        deck.save(source)
    else:
        source.write_text("{}", encoding="utf-8")
    return source, output


def task_prompt(app, output):
    tasks = {
        "notepad": (
            "In the open Notepad checklist, change the review date to 21 October 2026, "
            "change the status to Ready for review, and add '- Book the projector' as the "
            "last action. Keep all other text unchanged."
        ),
        "word": (
            "In the open Word document, format 'Cedar project update' with the Heading 1 style "
            "and change the launch date in the first paragraph to 21 October 2026. "
            "Keep the rest of the document's text and heading styles unchanged."
        ),
        "powerpoint": (
            "In the open PowerPoint presentation, change the first slide title to "
            "'Cedar launch briefing'. On the Schedule slide, change 'Launch: November' "
            "to 'Launch: December'. Move the Next steps slide immediately before Schedule. "
            "Keep all other slide content unchanged."
        ),
        "chrome": (
            "Complete and submit the workshop booking in the open Chrome page for Maya Chen "
            "(maya.chen@example.test), Operations department, 12 attendees, on 21 October 2026. "
            "Request a projector and enter the notes 'Bring printed agendas.'."
        ),
    }
    prompt = tasks[app]
    if app != "chrome":
        prompt += f" Save a copy to {output}, leaving the original file unchanged."
    return prompt


def verify_output(app, output):
    if not output.is_file():
        return {"success": False, "checks": {"saved_file_exists": False}}
    try:
        if app == "notepad":
            data = output.read_bytes()
            text = data.decode("utf-16" if data.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig")
            text = text.replace("\r\n", "\n")
            return {"success": text.rstrip("\n") == NOTEPAD_EXPECTED.rstrip("\n"), "text": text}
        if app == "word":
            from docx import Document
            document = Document(output)
            text = [p.text for p in document.paragraphs if p.text]
            styles = [p.style.name for p in document.paragraphs if p.text]
            checks = {
                "text": text == WORD_EXPECTED,
                "title_style": bool(styles) and styles[0] == "Heading 1",
                "other_styles": styles[1:] == ["Normal", "Heading 2", "Normal"],
            }
            return {"success": all(checks.values()), "checks": checks, "text": text, "styles": styles}
        if app == "powerpoint":
            from pptx import Presentation
            deck = Presentation(output)
            slides = [
                [p.text for shape in slide.shapes if shape.has_text_frame
                 for p in shape.text_frame.paragraphs if p.text]
                for slide in deck.slides
            ]
            return {"success": slides == POWERPOINT_EXPECTED, "slides": slides}
        values = json.loads(output.read_text(encoding="utf-8"))
        return {"success": values == CHROME_EXPECTED, "submitted": values}
    except (UnicodeError, zipfile.BadZipFile, ValueError, KeyError) as error:
        return {"success": False, "error": f"{type(error).__name__}: {error}"}
